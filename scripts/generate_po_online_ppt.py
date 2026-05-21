#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Generate a PO Online user guide PowerPoint.

Required packages:
    py -m pip install python-pptx openpyxl

Run:
    py scripts/generate_po_online_ppt.py
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

try:
    from openpyxl import Workbook, load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: openpyxl. Install with: py -m pip install openpyxl") from exc

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency: python-pptx. Install with: py -m pip install python-pptx") from exc

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = Path("storage/app/public/po-online-user-guide.pptx")
FONT_NAME = "Segoe UI"

BLUE = RGBColor(37, 99, 235)
BLUE_DARK = RGBColor(30, 64, 175)
TEXT = RGBColor(15, 23, 42)
MUTED = RGBColor(71, 85, 105)
LINE = RGBColor(203, 213, 225)
BG = RGBColor(248, 250, 252)
NOTE_FILL = RGBColor(239, 246, 255)
NOTE_LINE = RGBColor(147, 197, 253)
WARN_FILL = RGBColor(254, 249, 195)
WARN_LINE = RGBColor(250, 204, 21)

SLIDE_COLUMNS = [
    "slide_no",
    "layout",
    "section",
    "title",
    "subtitle",
    "steps",
    "bullets",
    "note",
    "warning",
    "image_path",
    "email_subject",
    "email_body",
    "notes",
]


@dataclass
class SlideRow:
    slide_no: int
    layout: str
    section: str
    title: str
    subtitle: str
    steps: List[str]
    bullets: List[str]
    note: str
    warning: str
    image_path: str
    email_subject: str
    email_body: str
    notes: str


TOC_ITEMS = [
    "ภาพรวมระบบ PO Online",
    "หน้ารายการ PO",
    "การเปิดดูรายละเอียด PO",
    "การแนบเอกสารประกอบ",
    "การส่ง PO เข้า Workflow",
    "การ Approve / Reject",
    "หน้า PO My Actions",
    "Workflow Comment / Remark",
    "การ Download PDF",
    "วิธีเปลี่ยนลายเซ็น",
    "ตัวอย่างเมล Approve",
    "ข้อควรจำ",
]


DEFAULT_ROWS = [
    {
        "layout": "cover",
        "section": "PO Online",
        "title": "คู่มือใช้งาน PO Online",
        "subtitle": "คู่มืออบรมพนักงานสำหรับการดู PO แนบเอกสาร ส่งอนุมัติ Approve / Reject Download PDF และเปลี่ยนลายเซ็น",
        "bullets": "ใช้ภาพหน้าจอจริงประกอบการทำงาน|อ่านตามลำดับแล้วทำตามได้ทันที|เหมาะสำหรับผู้ใช้งานทั่วไปและผู้อนุมัติ",
        "warning": "ผู้ใช้งานต้องมีสิทธิ์เข้าเมนู PO Online ก่อนเริ่มใช้งาน",
    },
    {
        "layout": "toc",
        "section": "สารบัญ",
        "title": "สารบัญ",
        "subtitle": "ลำดับหัวข้อในคู่มือนี้",
        "bullets": "|".join(TOC_ITEMS),
    },
    {
        "layout": "wide_image",
        "section": "ภาพรวมระบบ PO Online",
        "title": "ภาพรวมระบบ PO Online",
        "subtitle": "ระบบช่วยติดตาม PO ตั้งแต่แนบเอกสาร ส่งอนุมัติ จนถึง Download PDF หลังอนุมัติครบ",
        "steps": "เปิดหน้ารายการ PO|ตรวจสถานะและเอกสารแนบ|ทำงานตามปุ่ม Action ของแต่ละรายการ",
        "note": "ฝ่าย Purchase ทุกคนที่มีสิทธิ์ PO สามารถเป็นผู้ส่งเอกสารเข้า Workflow ได้",
        "warning": "ก่อนส่งเข้า Workflow ต้องมีเอกสารแนบเสมอ",
        "image_path": "screenshots/po-list.png",
    },
    {
        "layout": "two_col",
        "section": "หน้ารายการ PO",
        "title": "ค้นหาและกรองรายการ PO",
        "subtitle": "ใช้หน้ารายการเพื่อหา PO ที่ต้องตรวจหรือดำเนินการต่อ",
        "steps": "ค้นหาด้วย PO / Invoice / Vendor / แผนก|เลือกกลุ่มแผนก สถานะ Source หรือช่วงวันที่|กด เปิด เพื่อดูรายละเอียด PO",
        "note": "ถ้าต้องการกลับไปดูรายการทั้งหมด ให้กดล้างตัวกรอง",
        "warning": "หากหา PO ไม่เจอ ให้ตรวจตัวกรองก่อนเป็นอันดับแรก",
        "image_path": "screenshots/po-list.png",
    },
    {
        "layout": "two_col",
        "section": "การเปิดดูรายละเอียด PO",
        "title": "ตรวจรายละเอียด PO",
        "subtitle": "ตรวจข้อมูลหลักก่อนแนบเอกสารหรือส่งอนุมัติ",
        "steps": "เปิดรายการ PO|ตรวจ Vendor, Department และ Amount|ตรวจเอกสารแนบและประวัติด้านล่าง",
        "note": "ข้อมูลในหน้านี้ใช้ประกอบการตัดสินใจก่อนส่งต่อให้ผู้อนุมัติ",
        "warning": "ถ้าข้อมูล PO ไม่ถูกต้อง ให้ตรวจจาก ERP หรือแจ้งผู้เกี่ยวข้องก่อน",
        "image_path": "screenshots/po-detail.png",
    },
    {
        "layout": "two_col",
        "section": "การแนบเอกสารประกอบ",
        "title": "แนบเอกสารก่อนส่งอนุมัติ",
        "subtitle": "เอกสารแนบช่วยให้ผู้อนุมัติตรวจข้อมูลประกอบได้ครบ",
        "steps": "เปิดรายละเอียด PO|เลือกไฟล์เอกสารแล้วกดแนบไฟล์|ตรวจว่าสถานะเป็น Attached",
        "note": "แนบไฟล์ที่เกี่ยวข้องกับ PO เช่น ใบเสนอราคา หรือเอกสารประกอบการสั่งซื้อ",
        "warning": "ต้องแนบเอกสารก่อน ระบบจึงจะส่งเข้า Workflow ได้",
        "image_path": "screenshots/attachment.png",
    },
    {
        "layout": "two_col",
        "section": "การส่ง PO เข้า Workflow",
        "title": "การส่ง PO เข้าระบบอนุมัติ",
        "subtitle": "ใช้เมื่อเตรียมเอกสารครบและต้องการเริ่มขั้นตอนอนุมัติ",
        "steps": "กรอก Comment สำหรับส่งเข้า Workflow|กดปุ่ม ส่งเข้าระบบอนุมัติ|ระบบเปลี่ยนสถานะเป็น Waiting Approval",
        "note": "Comment จะถูกเก็บเป็นประวัติของเอกสาร",
        "warning": "ต้องมีเอกสารแนบก่อน จึงจะส่งเข้า Workflow ได้",
        "image_path": "screenshots/workflow-submit.png",
    },
    {
        "layout": "two_col",
        "section": "Workflow Approval",
        "title": "การตรวจสอบผู้ที่รออนุมัติ",
        "subtitle": "กล่อง Waiting Approval จะแสดงขั้นตอนและผู้ที่ต้องดำเนินการ",
        "bullets": "Current Step แสดงขั้นตอนปัจจุบัน|Purchase Approval คือผู้อนุมัติฝ่าย Purchase|Department Manager Approval คือหัวหน้าแผนกต้นทาง|ระบบแสดงชื่อ อีเมล และ Step ของผู้อนุมัติ",
        "note": "ดูส่วนนี้เมื่อต้องการรู้ว่าเอกสารค้างอยู่ที่ใคร",
        "warning": "ถ้าไม่เห็นผู้อนุมัติ แสดงว่า rule หรือข้อมูลตำแหน่งอาจยังไม่ครบ",
        "image_path": "screenshots/workflow-waiting-purchase.png",
    },
    {
        "layout": "two_col",
        "section": "หน้า PO My Actions",
        "title": "หน้า PO My Actions",
        "subtitle": "สำหรับดู PO ที่ผู้ใช้งานปัจจุบันมีสิทธิ์ดำเนินการ",
        "bullets": "ค้นหาด้วย PO / Invoice / Vendor / แผนก|กรองตาม Status, Source และช่วงวันที่|กด เปิด เพื่อเข้าไป Approve หรือ Reject",
        "note": "หน้านี้จะแสดงเฉพาะงานที่เกี่ยวข้องกับผู้ใช้งานปัจจุบัน",
        "warning": "ถ้าไม่มีรายการ แปลว่าไม่มี PO ที่ถึง Step ของผู้ใช้งานคนนี้",
        "image_path": "screenshots/my-actions.png",
    },
    {
        "layout": "two_col",
        "section": "การ Approve / Reject",
        "title": "วิธี Approve เอกสาร",
        "subtitle": "ผู้อนุมัติเปิดเอกสารจาก My Actions แล้วเลือก Approve หรือ Reject",
        "steps": "เปิดเอกสารจากหน้า PO My Actions|กรอก Approve Comment|กด Approve เพื่อส่งต่อ Step ถัดไป",
        "note": "ถ้าต้องการส่งกลับหรือไม่อนุมัติ ให้กด Send Back / Reject",
        "warning": "ควรตรวจเอกสารแนบและรายละเอียด PO ก่อนกด Approve ทุกครั้ง",
        "image_path": "screenshots/workflow-purchase-approve.png",
    },
    {
        "layout": "two_col",
        "section": "การ Approve / Reject",
        "title": "รอ Department Manager Approval",
        "subtitle": "หลังฝ่าย Purchase อนุมัติครบ ระบบจะส่งต่อไปยังหัวหน้าแผนกต้นทาง",
        "bullets": "สถานะจะแสดงว่าอยู่ที่ Department Manager Approval|ผู้มีสิทธิ์จะเห็นรายการใน PO My Actions|หัวหน้าแผนกสามารถ Approve หรือ Reject ได้จากหน้ารายละเอียด",
        "note": "จัดซื้อสามารถเลือกส่ง Mail แจ้งหัวหน้าแผนกได้เอง",
        "warning": "รายการที่ยังไม่ถึง Step นี้จะยังไม่สามารถส่ง Mail แจ้งหัวหน้าแผนกได้",
        "image_path": "screenshots/workflow-waiting-manager.png",
    },
    {
        "layout": "two_col",
        "section": "การ Approve / Reject",
        "title": "ผู้จัดการกด Approve / Reject",
        "subtitle": "หน้าจออนุมัติของหัวหน้าแผนกใช้หลักการเดียวกับฝ่าย Purchase",
        "steps": "เปิดเอกสารที่รอดำเนินการ|กรอก Comment|กด Approve หรือ Reject",
        "note": "เมื่ออนุมัติครบทุก Step ระบบจะแสดง Process completed",
        "warning": "Comment ช่วยให้ตรวจสอบย้อนหลังได้ง่าย ควรกรอกทุกครั้ง",
        "image_path": "screenshots/workflow-manager-approve.png",
    },
    {
        "layout": "two_col",
        "section": "Workflow Comment / Remark",
        "title": "Workflow Comment / Remark",
        "subtitle": "แสดงประวัติการดำเนินการทั้งหมดของเอกสาร",
        "bullets": "SUBMIT = ผู้ส่งเอกสารเข้าระบบ|APPROVE = ผู้อนุมัติแต่ละ Step|COMPLETE = เอกสารอนุมัติครบทุกขั้นตอน|มีวันที่ เวลา ผู้ดำเนินการ และ Comment",
        "note": "ใช้ส่วนนี้เมื่อต้องการย้อนดูว่าใครทำอะไรและเมื่อไร",
        "warning": "ถ้าเอกสารยังไม่ถูกส่งเข้า Workflow จะยังไม่มีประวัติ Action",
        "image_path": "screenshots/workflow-complete.png",
    },
    {
        "layout": "flow",
        "section": "Flow การอนุมัติ PO Online",
        "title": "Flow การอนุมัติ PO Online",
        "subtitle": "ลำดับการทำงานหลักของ Workflow",
        "bullets": "Submit PO|Purchase Approval|Department Manager Approval|Complete",
        "note": "หลังอนุมัติครบทุก Step ระบบจะแสดง Process completed",
        "warning": "ผู้ใช้งานจะเห็นปุ่ม Approve เฉพาะเอกสารที่ถึง Step ของตัวเอง",
    },
    {
        "layout": "two_col",
        "section": "การ Download PDF",
        "title": "Download PDF ของ PO",
        "subtitle": "ระบบดึง PDF จาก ERP และแปะลายเซ็นผู้อนุมัติให้ตามขั้นตอน",
        "steps": "เลือกรายการ PO ที่ต้องการ|กด Download PDF รายการเดียวหรือทั้งกลุ่ม|เปิด PDF เพื่อตรวจลายเซ็นและวันที่",
        "note": "PDF ที่ได้อ้างอิงแบบฟอร์มจาก ERP",
        "warning": "ถ้า Download ไม่สำเร็จ ให้เชื่อม ERP ใหม่ หรือแจ้งผู้ดูแลระบบ",
        "image_path": "screenshots/pdf-example.png",
    },
    {
        "layout": "two_col",
        "section": "วิธีเปลี่ยนลายเซ็น",
        "title": "เปลี่ยนลายเซ็นใน Profile",
        "subtitle": "ลายเซ็นนี้จะถูกใช้ใน PDF หลังจากอนุมัติเอกสาร",
        "steps": "เปิดเมนู Profile ของตัวเอง|อัปโหลดไฟล์ลายเซ็นใหม่|กดบันทึก แล้วเปิด PDF เพื่อตรวจผล",
        "note": "แนะนำให้ใช้ PNG หรือ JPG พื้นหลังขาว ลายเซ็นชัด",
        "warning": "อย่าใช้ลายเซ็นของผู้อื่น และควรเปลี่ยนเฉพาะบัญชีของตัวเอง",
        "image_path": "screenshots/profile.png",
    },
    {
        "layout": "email",
        "section": "ตัวอย่างเมล Approve",
        "title": "ตัวอย่างอีเมล Approve",
        "subtitle": "อีเมลแจ้งให้ผู้อนุมัติเข้าไปตรวจและอนุมัติใน PO Online",
        "steps": "เปิดอีเมลที่ได้รับ|ตรวจเลข PO และข้อมูลสำคัญ|คลิกลิงก์เพื่อเปิดเอกสารใน PO Online",
        "note": "ผู้อนุมัติควรเปิดเอกสารจริงในระบบก่อนตัดสินใจ",
        "warning": "ไม่ควรอนุมัติจากข้อมูลในอีเมลอย่างเดียว",
        "email_subject": "PO Online Approval Request - PO No. {po_no}",
        "email_body": "\n".join(
            [
                "Dear Approver,",
                "Please review and approve the purchase order in PO Online.",
                "",
                "PO No.: {po_no}",
                "Vendor: {vendor}",
                "Department: {department}",
                "Amount: {amount} THB",
                "",
                "Please click the link below to open and approve the document.",
                "Thank you.",
            ]
        ),
    },
    {
        "layout": "summary",
        "section": "ข้อควรจำ",
        "title": "ข้อควรจำสำหรับ Workflow",
        "subtitle": "สรุปสิ่งที่ควรตรวจทุกครั้ง",
        "bullets": "ถ้าเอกสารยังไม่ถูกส่งเข้า Workflow จะขึ้นว่ายังไม่มี Action|ผู้ใช้งานจะเห็นปุ่ม Approve เฉพาะเอกสารที่ถึง Step ของตัวเอง|หลังอนุมัติครบทุก Step ระบบจะแสดง Process completed|ควรกรอก Comment ทุกครั้งเพื่อให้ตรวจสอบย้อนหลังได้",
        "note": "เมื่อไม่แน่ใจว่าเอกสารอยู่ที่ใคร ให้ดู Current Step และรายชื่อ Pending Approver",
        "warning": "อย่ากด Approve หากยังไม่ได้ตรวจเอกสารแนบและข้อมูล PO",
    },
]


def split_lines(value: object) -> List[str]:
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    if "\n" not in text and "|" in text:
        text = text.replace("|", "\n")
    return [line.strip(" -\t") for line in text.split("\n") if line.strip()]


def cell_text(value: object) -> str:
    return str(value or "").strip()


def default_slide_rows() -> List[SlideRow]:
    rows: List[SlideRow] = []
    for index, row in enumerate(DEFAULT_ROWS, start=1):
        rows.append(
            SlideRow(
                slide_no=index,
                layout=cell_text(row.get("layout", "two_col")).lower(),
                section=cell_text(row.get("section")),
                title=cell_text(row.get("title")),
                subtitle=cell_text(row.get("subtitle")),
                steps=split_lines(row.get("steps")),
                bullets=split_lines(row.get("bullets")),
                note=cell_text(row.get("note")),
                warning=cell_text(row.get("warning")),
                image_path=cell_text(row.get("image_path")),
                email_subject=cell_text(row.get("email_subject")),
                email_body=str(row.get("email_body", "") or "").strip(),
                notes=cell_text(row.get("notes")),
            )
        )
    return rows


def load_rows(excel_path: Path) -> List[SlideRow]:
    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook["Slides"] if "Slides" in workbook.sheetnames else workbook.active
    header = [cell_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    index = {name: header.index(name) for name in SLIDE_COLUMNS if name in header}
    missing = [name for name in ["layout", "title"] if name not in index]
    if missing:
        raise SystemExit(f"Excel is missing required column(s): {', '.join(missing)}")

    rows: List[SlideRow] = []
    for cells in sheet.iter_rows(min_row=2):
        values = [cell.value for cell in cells]

        def value(column: str, default: object = "") -> object:
            if column not in index or index[column] >= len(values):
                return default
            return values[index[column]] if values[index[column]] is not None else default

        title = cell_text(value("title"))
        if not title:
            continue
        rows.append(
            SlideRow(
                slide_no=int(value("slide_no", len(rows) + 1) or len(rows) + 1),
                layout=cell_text(value("layout", "two_col")).lower(),
                section=cell_text(value("section")),
                title=title,
                subtitle=cell_text(value("subtitle")),
                steps=split_lines(value("steps")),
                bullets=split_lines(value("bullets")),
                note=cell_text(value("note")),
                warning=cell_text(value("warning")),
                image_path=cell_text(value("image_path")),
                email_subject=cell_text(value("email_subject")),
                email_body=str(value("email_body", "") or "").strip(),
                notes=cell_text(value("notes")),
            )
        )
    return sorted(rows, key=lambda item: item.slide_no)


def create_template(output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Slides"
    sheet.append(SLIDE_COLUMNS)
    for row in default_slide_rows():
        sheet.append(
            [
                row.slide_no,
                row.layout,
                row.section,
                row.title,
                row.subtitle,
                "\n".join(row.steps),
                "\n".join(row.bullets),
                row.note,
                row.warning,
                row.image_path,
                row.email_subject,
                row.email_body,
                row.notes,
            ]
        )
    for col in "ABCDEFGHIJKLM":
        sheet.column_dimensions[col].width = 28 if col not in ["F", "G", "H", "I", "L"] else 58
    workbook.save(output_path)


def text_box(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    size: int,
    color=TEXT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = FONT_NAME
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return shape


def set_fill(shape, fill, line=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line


def header(slide, row: SlideRow) -> None:
    text_box(slide, Inches(0.55), Inches(0.25), Inches(8.0), Inches(0.32), row.section, 12, BLUE, True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.66), Inches(12.2), Inches(0.02))
    set_fill(rule, RGBColor(226, 232, 240))


def title_block(slide, row: SlideRow, width=Inches(5.2)) -> None:
    text_box(slide, Inches(0.65), Inches(0.88), width, Inches(0.55), row.title, 34, TEXT, True)
    if row.subtitle:
        text_box(slide, Inches(0.68), Inches(1.48), width, Inches(0.7), row.subtitle, 22, MUTED)


def icon_label(slide, text: str, left, top) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(0.48), Inches(0.36))
    set_fill(box, BLUE, BLUE)
    frame = box.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.text = text[:4].upper()
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)


def steps(slide, items: Sequence[str], left, top, width, size=18, max_items=4) -> None:
    for index, item in enumerate(list(items)[:max_items], start=1):
        row_top = top + Inches((index - 1) * 0.72)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, row_top, Inches(0.43), Inches(0.43))
        set_fill(circle, BLUE, BLUE)
        frame = circle.text_frame
        frame.clear()
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]
        p.text = str(index)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        text_box(slide, left + Inches(0.58), row_top - Inches(0.02), width - Inches(0.58), Inches(0.5), item, size, TEXT, True)


def bullets(slide, items: Iterable[str], left, top, width, height, size=18) -> None:
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(list(items)[:5]):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT_NAME
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(7)


def info_box(slide, label: str, body: str, left, top, width, height, kind: str) -> None:
    if not body:
        return
    fill = WARN_FILL if kind == "warning" else NOTE_FILL
    line = WARN_LINE if kind == "warning" else NOTE_LINE
    label_color = RGBColor(133, 77, 14) if kind == "warning" else BLUE_DARK
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_fill(box, fill, line)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.17)
    frame.margin_right = Inches(0.12)
    frame.margin_top = Inches(0.08)
    p = frame.paragraphs[0]
    p.text = label
    p.font.name = FONT_NAME
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = label_color
    p.space_after = Pt(4)
    p = frame.add_paragraph()
    p.text = body
    p.font.name = FONT_NAME
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT


def resolve_image_path(image_path: str, base_dir: Path) -> Optional[Path]:
    if not image_path:
        return None
    raw = Path(image_path)
    candidates = [raw] if raw.is_absolute() else [base_dir / raw, Path.cwd() / raw, REPO_ROOT / raw]
    return next((candidate for candidate in candidates if candidate.exists()), None)


def contained_size(path: Path, max_width, max_height):
    if Image is None:
        return max_width, max_height
    with Image.open(path) as img:
        width_px, height_px = img.size
    ratio = min(max_width / width_px, max_height / height_px)
    return int(width_px * ratio), int(height_px * ratio)


def image_frame(slide, image_path: str, base_dir: Path, left, top, width, height, missing_images: List[str]) -> None:
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_fill(frame, RGBColor(255, 255, 255), LINE)
    path = resolve_image_path(image_path, base_dir)
    if path:
        img_w, img_h = contained_size(path, width - Inches(0.18), height - Inches(0.18))
        img_left = left + int((width - img_w) / 2)
        img_top = top + int((height - img_h) / 2)
        slide.shapes.add_picture(str(path), img_left, img_top, width=img_w, height=img_h)
        return
    if image_path:
        missing_images.append(image_path)
    text_box(slide, left + Inches(0.25), top + Inches(1.6), width - Inches(0.5), Inches(0.8), f"วางรูปไว้ที่ {image_path}", 20, MUTED, True, PP_ALIGN.CENTER)


def cover_slide(slide, row: SlideRow) -> None:
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.22), Inches(7.5))
    set_fill(band, BLUE, BLUE)
    icon_label(slide, "PO", Inches(0.82), Inches(0.85))
    text_box(slide, Inches(0.82), Inches(1.28), Inches(10.8), Inches(0.85), row.title, 36, TEXT, True)
    text_box(slide, Inches(0.86), Inches(2.18), Inches(10.7), Inches(0.7), row.subtitle, 24, MUTED)
    steps(slide, row.bullets, Inches(1.0), Inches(3.35), Inches(9.8), 20, 3)
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(0.9), Inches(6.15), Inches(10.8), Inches(0.72), "warning")


def toc_slide(slide, row: SlideRow) -> None:
    header(slide, row)
    title_block(slide, row, Inches(11.8))
    items = list(row.bullets)
    for i, item in enumerate(items):
        col = 0 if i < 6 else 1
        pos = i if i < 6 else i - 6
        left = Inches(0.9 + (col * 6.0))
        top = Inches(2.35 + pos * 0.58)
        icon_label(slide, str(i + 1), left, top)
        text_box(slide, left + Inches(0.62), top - Inches(0.02), Inches(4.9), Inches(0.42), item, 20, TEXT, True)


def two_col_slide(slide, row: SlideRow, base_dir: Path, missing_images: List[str]) -> None:
    header(slide, row)
    title_block(slide, row)
    if row.steps:
        steps(slide, row.steps, Inches(0.78), Inches(2.32), Inches(4.95), 18, 4)
        bullet_top = Inches(5.10)
    else:
        bullets(slide, row.bullets, Inches(0.78), Inches(2.28), Inches(4.95), Inches(2.65), 18)
        bullet_top = Inches(5.10)
    if row.steps and row.bullets:
        bullets(slide, row.bullets, Inches(0.8), Inches(4.88), Inches(4.9), Inches(0.8), 15)
    info_box(slide, "Note", row.note, Inches(0.68), bullet_top, Inches(5.15), Inches(0.72), "note")
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(0.68), Inches(6.02), Inches(5.15), Inches(0.80), "warning")
    image_frame(slide, row.image_path, base_dir, Inches(6.05), Inches(0.98), Inches(6.58), Inches(5.9), missing_images)


def wide_image_slide(slide, row: SlideRow, base_dir: Path, missing_images: List[str]) -> None:
    header(slide, row)
    text_box(slide, Inches(0.65), Inches(0.88), Inches(11.8), Inches(0.5), row.title, 34, TEXT, True)
    text_box(slide, Inches(0.68), Inches(1.42), Inches(11.4), Inches(0.48), row.subtitle, 22, MUTED)
    steps(slide, row.steps, Inches(0.78), Inches(2.05), Inches(11.4), 18, 3)
    image_frame(slide, row.image_path, base_dir, Inches(0.72), Inches(4.25), Inches(7.95), Inches(2.55), missing_images)
    info_box(slide, "Note", row.note, Inches(8.95), Inches(4.32), Inches(3.55), Inches(0.88), "note")
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(8.95), Inches(5.72), Inches(3.55), Inches(0.88), "warning")


def flow_slide(slide, row: SlideRow) -> None:
    header(slide, row)
    text_box(slide, Inches(0.65), Inches(0.88), Inches(11.6), Inches(0.55), row.title, 36, TEXT, True)
    text_box(slide, Inches(0.68), Inches(1.48), Inches(11.0), Inches(0.45), row.subtitle, 24, MUTED)
    x_positions = [1.0, 4.0, 7.1, 10.2]
    labels = list(row.bullets)[:4]
    for i, label in enumerate(labels):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x_positions[i]), Inches(3.05), Inches(2.2), Inches(0.92))
        set_fill(box, BLUE if i < 3 else RGBColor(22, 163, 74), BLUE)
        frame = box.text_frame
        frame.clear()
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        if i < 3:
            text_box(slide, Inches(x_positions[i] + 2.28), Inches(3.22), Inches(0.55), Inches(0.35), ">", 26, BLUE_DARK, True, PP_ALIGN.CENTER)
    info_box(slide, "Note", row.note, Inches(1.0), Inches(5.2), Inches(5.4), Inches(0.88), "note")
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(6.9), Inches(5.2), Inches(4.9), Inches(0.88), "warning")


def email_slide(slide, row: SlideRow) -> None:
    header(slide, row)
    title_block(slide, row)
    steps(slide, row.steps, Inches(0.82), Inches(2.3), Inches(4.8), 18, 3)
    info_box(slide, "Note", row.note, Inches(0.68), Inches(5.1), Inches(5.15), Inches(0.72), "note")
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(0.68), Inches(6.08), Inches(5.15), Inches(0.72), "warning")
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.1), Inches(2.0), Inches(6.25), Inches(4.38))
    set_fill(card, RGBColor(255, 255, 255), LINE)
    text_box(slide, Inches(6.35), Inches(2.28), Inches(5.7), Inches(0.32), "Subject:", 15, MUTED, True)
    text_box(slide, Inches(6.35), Inches(2.64), Inches(5.7), Inches(0.42), row.email_subject, 18, TEXT, True)
    text_box(slide, Inches(6.35), Inches(3.22), Inches(5.7), Inches(2.75), row.email_body, 18, TEXT)


def summary_slide(slide, row: SlideRow) -> None:
    header(slide, row)
    title_block(slide, row, Inches(11.8))
    items = list(row.bullets)
    for i, item in enumerate(items[:4], start=1):
        top = Inches(2.22 + (i - 1) * 0.76)
        steps(slide, [item], Inches(1.0), top, Inches(10.6), 20, 1)
    info_box(slide, "Note", row.note, Inches(1.0), Inches(5.55), Inches(5.3), Inches(0.82), "note")
    info_box(slide, "ข้อควรระวัง", row.warning, Inches(6.75), Inches(5.55), Inches(5.1), Inches(0.82), "warning")


def add_notes(slide, notes: str) -> None:
    if not notes:
        return
    try:
        slide.notes_slide.notes_text_frame.text = notes
    except Exception:
        return


def build_deck(rows: List[SlideRow], output_path: Path, base_dir: Path) -> List[str]:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    missing_images: List[str] = []

    for row in rows:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG
        if row.layout == "cover":
            cover_slide(slide, row)
        elif row.layout == "toc":
            toc_slide(slide, row)
        elif row.layout == "wide_image":
            wide_image_slide(slide, row, base_dir, missing_images)
        elif row.layout == "flow":
            flow_slide(slide, row)
        elif row.layout == "email":
            email_slide(slide, row)
        elif row.layout == "summary":
            summary_slide(slide, row)
        else:
            two_col_slide(slide, row, base_dir, missing_images)
        add_notes(slide, row.notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return sorted(set(missing_images))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate PO Online user guide PowerPoint.")
    parser.add_argument("--excel", type=Path, help="Input Excel workbook with a Slides sheet.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output PPTX path.")
    parser.add_argument("--init-template", type=Path, help="Create an Excel template and exit.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.init_template:
        create_template(args.init_template)
        print(f"Created Excel template: {args.init_template}")
        return

    if args.excel:
        rows = load_rows(args.excel)
        base_dir = args.excel.parent
    else:
        rows = default_slide_rows()
        base_dir = REPO_ROOT

    if not rows:
        raise SystemExit("No slide rows found.")

    missing_images = build_deck(rows, args.output, base_dir)
    print(f"Created PowerPoint: {args.output}")
    if missing_images:
        print("Missing screenshot file(s):")
        for image in missing_images:
            print(f" - {image}")


if __name__ == "__main__":
    main()
